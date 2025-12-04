#!/usr/bin/env python3
"""
Script para converter CRONOGRAMA DE ESTUDOS EA.xlsx em HTML completo
"""

import openpyxl
import json

def converter_excel_para_json(arquivo_excel):
    """Converte o Excel em estrutura JSON"""

    wb = openpyxl.load_workbook(arquivo_excel)
    semanas = []

    # Para cada planilha de semana (SEMANA 1 até SEMANA 24)
    for num_semana in range(1, 25):
        nome_sheet = f"SEMANA {num_semana}"

        if nome_sheet not in wb.sheetnames:
            print(f"⚠️  Planilha '{nome_sheet}' não encontrada. Pulando...")
            continue

        sheet = wb[nome_sheet]
        dias = []

        # Assumindo que os dados começam na linha 2 (linha 1 é cabeçalho)
        # Colunas: A=DIA, B=DISCIPLINA, C=TEMA, D=TEORIA, E=EXERCÍCIOS

        for row in sheet.iter_rows(min_row=2, values_only=True):
            dia = row[0]  # Coluna A
            disciplina = row[1]  # Coluna B
            tema = row[2]  # Coluna C
            teoria = row[3] if len(row) > 3 else ""  # Coluna D
            exercicios = row[4] if len(row) > 4 else ""  # Coluna E

            # Pular linhas vazias
            if not dia or not disciplina:
                continue

            dias.append({
                "dia": int(dia) if isinstance(dia, (int, float)) else dia,
                "disciplina": str(disciplina).strip(),
                "tema": str(tema).strip() if tema else "",
                "teoria": str(teoria).strip() if teoria else "",
                "exercicios": str(exercicios).strip() if exercicios else "",
                "concluido": False,
                "tempoEstudado": 0
            })

        # Verificar se é semana de revisão (5, 10, 15, 20-24)
        eh_revisao = num_semana in [5, 10, 15, 20, 21, 22, 23, 24]

        semanas.append({
            "id": num_semana,
            "nome": f"SEMANA {num_semana}",
            "revisao": eh_revisao,
            "dias": dias
        })

        print(f"✅ Semana {num_semana}: {len(dias)} atividades")

    return {"semanas": semanas}


def gerar_html_completo(dados_json, arquivo_saida):
    """Gera o arquivo HTML completo com os dados"""

    # Ler o template HTML
    with open('cronograma-wordpress-FINAL.html', 'r', encoding='utf-8') as f:
        html_template = f.read()

    # Converter dados para JavaScript
    dados_js = json.dumps(dados_json, ensure_ascii=False, indent=4)

    # Substituir o placeholder no template
    # Procurar por: const cronogramaDataDefault = {
    inicio = html_template.find('const cronogramaDataDefault = {')

    if inicio == -1:
        print("❌ Erro: Não encontrei 'const cronogramaDataDefault' no template")
        return False

    # Encontrar o final do objeto (procurar por '};' depois do início)
    fim = html_template.find('};', inicio)

    if fim == -1:
        print("❌ Erro: Não encontrei o final do objeto cronogramaDataDefault")
        return False

    # Substituir
    novo_html = (
        html_template[:inicio] +
        f'const cronogramaDataDefault = {dados_js}' +
        html_template[fim+1:]
    )

    # Salvar
    with open(arquivo_saida, 'w', encoding='utf-8') as f:
        f.write(novo_html)

    return True


if __name__ == "__main__":
    import sys

    print("=" * 60)
    print("🚀 CONVERSOR EXCEL → HTML COMPLETO")
    print("=" * 60)
    print()

    # Verificar argumentos
    if len(sys.argv) < 2:
        print("📖 USO:")
        print(f"   python {sys.argv[0]} ARQUIVO_EXCEL.xlsx")
        print()
        print("📝 EXEMPLO:")
        print(f"   python {sys.argv[0]} 'CRONOGRAMA DE ESTUDOS EA.xlsx'")
        print()
        sys.exit(1)

    arquivo_excel = sys.argv[1]

    try:
        # 1. Converter Excel para JSON
        print("📂 Lendo arquivo Excel...")
        dados = converter_excel_para_json(arquivo_excel)

        total_atividades = sum(len(s['dias']) for s in dados['semanas'])
        print()
        print(f"✅ Convertido com sucesso!")
        print(f"   • {len(dados['semanas'])} semanas")
        print(f"   • {total_atividades} atividades no total")
        print()

        # 2. Gerar HTML completo
        print("🔨 Gerando HTML completo...")
        arquivo_saida = "cronograma-completo-FINAL.html"

        if gerar_html_completo(dados, arquivo_saida):
            print(f"✅ HTML gerado: {arquivo_saida}")
            print()
            print("=" * 60)
            print("🎉 PRONTO!")
            print("=" * 60)
            print()
            print(f"📄 Abra o arquivo '{arquivo_saida}' no navegador para testar")
            print(f"📋 Copie TODO o conteúdo do arquivo para o WordPress")
            print()
        else:
            print("❌ Erro ao gerar HTML")
            sys.exit(1)

    except FileNotFoundError:
        print(f"❌ Erro: Arquivo '{arquivo_excel}' não encontrado")
        sys.exit(1)
    except Exception as e:
        print(f"❌ Erro: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)
